import csv, json, re, sys, os
S=str(__import__('pathlib').Path(__file__).resolve().parents[2] / 'data' / 'raw')
P=S+'/programs'; R=P+'/sources'
def w(name, rows, extra=()):
    cols=['year','nominal_usd_millions','note']+list(extra)
    with open(f'{P}/{name}.csv','w',newline='') as f:
        wr=csv.writer(f); wr.writerow(cols)
        for r in rows: wr.writerow([r.get(c,'') for c in cols])
    tot=sum(float(r['nominal_usd_millions']) for r in rows if r.get('nominal_usd_millions') not in ('',None))
    print(f'{name}: {len(rows)} rows, nominal total {tot:,.1f}')
    return tot
def fmt(x,nd=1): return '' if x is None else f'{x:.{nd}f}'

# ---------- MANHATTAN (Hewlett & Anderson 1962, App. 2, monthly expenditures, calendar years)
man={1942:16122,1943:344587,1944:939395,1945:610311,1946:281043}
w('manhattan',[dict(year=y,nominal_usd_millions=fmt(v/1000,3),note=('Aug-Dec 1942 only; includes $15.0M transferred to OSRD in Aug 1942' if y==1942 else 'calendar year')+'; MED expenditures, Hewlett & Anderson (1962) App. 2') for y,v in man.items()])

# ---------- APOLLO (Planetary Society Cost of Apollo, 'Apollo Program (original $)', FY obligations)
import openpyxl
wb=openpyxl.load_workbook(R+'/tps_cost_of_apollo.xlsx',data_only=True)
ws=wb['Apollo Program (original $)']
rows=[]
for r in ws.iter_rows(min_row=3,values_only=True):
    if isinstance(r[0],(int,float)) and 1960<=r[0]<=1973:
        direct,indirect=r[4],r[20]
        rows.append(dict(year=int(r[0]),nominal_usd_millions=fmt((direct+indirect)/1000,3),
            note=f'NASA fiscal year obligations; Project Apollo = direct R&D {direct/1000:.1f} + indirect (CoF, ops/overhead, tracking) {indirect/1000:.1f}; excludes Gemini & robotic lunar',
            apollo_direct_usd_millions=fmt(direct/1000,3), lunar_effort_total_usd_millions=fmt(r[3]/1000,3)))
w('apollo',rows,extra=('apollo_direct_usd_millions','lunar_effort_total_usd_millions'))

# ---------- INTERSTATE (FHWA Highway Statistics Summary to 1995, Table FA-203, Interstate column, thousands $)
rows=[]
for l in open(P+'/work/fa203.txt'):
    t=l.split()
    if t and re.fullmatch(r'19[4-9]\d',t[0]):
        y=int(t[0]);
        if t[1]=='-' : continue
        v=float(t[1].replace(',',''))
        if 1956<=y<=1995:
            rows.append(dict(year=y,nominal_usd_millions=fmt(v/1000,3),
              note='federal-aid Interstate expenditure of federal funds (federal share only; state match ~10% not included); FY basis per FHWA'+('; post-ISTEA: Interstate construction only (Interstate Maintenance/Substitution reported separately)' if y>=1992 else ''),
              total_incl_state_match_est_usd_millions=fmt(v/1000*128.9/114.3,3)))
        elif y<1956: print('pre-1956 interstate',y,v)
w('interstate',rows,extra=('total_incl_state_match_est_usd_millions',))

# ---------- MARSHALL (USAID Greenbook obligations, ECA/MSA grants+loans to ERP participants, FY)
import pandas as pd
df=pd.read_excel(R+'/usaid_greenbook_us_foreignaid_greenbook_heiss_mirror.xlsx',header=6)
df['FY']=pd.to_numeric(df['Fiscal Year'],errors='coerce')
d=df[df.FY.between(1949,1952)&df['Funding Account Name'].isin(['INACTIVE - USAID Grants','INACTIVE - USAID Loans'])]
erp=d[(d.Region.str.contains('Europe')&~d.Country.isin(['Serbia and Montenegro, Former','Spain']))|((d.Country=='Indonesia')&d.FY.isin([1949,1950]))]
s=erp.groupby('FY')['Obligations (Historical Dollars)'].sum()/1e6
lab={1949:'FY1949 = 3 Apr 1948-30 Jun 1949 (15 months)',1950:'FY1950 = Jul 1949-Jun 1950',1951:'FY1951 = Jul 1950-Jun 1951',1952:'FY1952 = Jul 1951-Jun 1952 (ERP wind-down under MSA)'}
crs={1949:6220.6,1950:4060.2,1951:2254.1}
w('marshall',[dict(year=int(y),nominal_usd_millions=fmt(v,1),note=lab[int(y)]+'; USAID Greenbook obligations, ECA/MSA grants+loans to 16 ERP countries + regional + Indonesia (Dutch East Indies, FY49-50)',
   crs_funds_made_available_usd_millions=fmt(crs.get(int(y))) ) for y,v in s.items()],extra=('crs_funds_made_available_usd_millions',))

# ---------- F-35: no annual data obtained
w('f35',[])

# ---------- ISS (TPS Historical NASA Budget workbook FY sheets, 'Actual' column)
wb=openpyxl.load_workbook(R+'/tps_historical_nasa_budget.xlsx',read_only=True,data_only=True)
rows=[]
for ws in wb:
    m=re.match(r'FY ?(\d{4})$',ws.title)
    if not m: continue
    y=int(m.group(1)); rr=[list(r) for r in ws.iter_rows(values_only=True)]
    hdr=rr[0]; ac=[i for i,h in enumerate(hdr) if str(h).strip().lower()=='actual'][0]
    get=lambda names: next((r[ac] for r in rr[1:] if r and r[0] in names and ac<len(r) and r[ac] is not None),None)
    iss=get(('ISS','International Space Station')); st=get(('Space Transportation',))
    if iss is None: continue
    if y<=2016:
        rows.append(dict(year=y,nominal_usd_millions=fmt(iss,1),note='NASA FY actual, "ISS" line (incl. ISS crew & cargo transportation through FY2016)',iss_line_usd_millions=fmt(iss,1),space_transportation_usd_millions=''))
    else:
        rows.append(dict(year=y,nominal_usd_millions=fmt(iss+st,1),note='NASA FY actual, ISS line + Space Transportation line (split from ISS in FY2017; Space Transportation also includes Commercial Crew development)',iss_line_usd_millions=fmt(iss,1),space_transportation_usd_millions=fmt(st,1)))
for r in rows:
    if r['year']==2026: r['note']+='; FY2026 "actual" as recorded in TPS FY2026 sheet (from FY2027 budget request; treat as preliminary)'
rows.sort(key=lambda r:r['year'])
w('iss',rows,extra=('iss_line_usd_millions','space_transportation_usd_millions'))

# ---------- RAILROADS
u=json.load(open(P+'/work/ulmer.json')); g=json.load(open(P+'/work/gallman.json'))
miles={}
for l in open(R+'/nber_macrohistory_a02082a.dat'):
    t=l.split()
    if len(t)>=2: miles[int(t[0])]=float(t[1])
pre60=990.7; book60_unadj=1151.6; book70=2200.0
m1=sum(miles[y] for y in range(1830,1861)); m2=sum(miles[y] for y in range(1861,1870))
rows=[]
for y in range(1830,1910):
    r=dict(year=y)
    if y<=1860:
        r['nominal_usd_millions']=fmt(pre60*miles[y]/m1,2)
        r['note']=f'DERIVED: Fishlow-adjusted cumulative construction cost through 1860 ($990.7M, cost basis, excl. land) distributed by miles built ({miles[y]:.0f} mi)'
    elif y<=1869:
        r['nominal_usd_millions']=fmt((book70-book60_unadj)*miles[y]/m2,2)
        r['note']=f'DERIVED: change in book cost of road & equipment 1860->1 Jan 1870 ($2,200M - $1,151.6M = $1,048.4M, unadjusted book basis) distributed by miles built ({miles[y]:.0f} mi)'
    else:
        c=u['c10'][str(y)]
        r['nominal_usd_millions']=fmt(c['excl_land'],0)
        r['note']=f"Ulmer (1960) Table C-10 col.3 gross capital expenditures excl. land, calendar year (incl. land: {c['incl_land']:.0f})"
    r['miles_built']=fmt(miles[y],0)
    if str(y) in g['t1']:
        t=g['t1'][str(y)]; r['gallman_rr_construction_1860usd_millions']=fmt(t['rr']); r['gallman_gnp_1860usd_millions']=fmt(t['gnp'])
        r['gallman_rr_construction_original_1860usd_millions']=fmt(t['rr'])
        r['gallman_year_basis']='census year ending 31 May of following year'
    if str(y) in g['t2']:
        t=g['t2'][str(y)]
        r['gallman_rr_construction_1860usd_millions']=fmt(t.get('rr_rev')) if 'rr_rev' in t else ''
        r['gallman_rr_construction_original_1860usd_millions']=fmt(t['rr_orig'])
        r['gallman_gnp_1860usd_millions']=fmt(t['gnp'])
        r['gallman_year_basis']='calendar year'+('; orig RR & GNP use Rhode decimal-error corrections' if 'gnp_uncorrected' in t else '')
    if r.get('gallman_rr_construction_1860usd_millions') and r.get('gallman_gnp_1860usd_millions'):
        r['gallman_rr_share_of_gnp_pct_1860prices']=fmt(100*float(r['gallman_rr_construction_1860usd_millions'])/float(r['gallman_gnp_1860usd_millions']),3)
    rows.append(r)
extra=('miles_built','gallman_rr_construction_1860usd_millions','gallman_gnp_1860usd_millions','gallman_rr_share_of_gnp_pct_1860prices','gallman_rr_construction_original_1860usd_millions','gallman_year_basis')
w('railroads',rows,extra=extra)
print('railroads 1830-1900 nominal:', sum(float(r['nominal_usd_millions']) for r in rows if r['year']<=1900))
print('  pre1861',sum(float(r['nominal_usd_millions']) for r in rows if r['year']<=1860),' 1861-69',sum(float(r['nominal_usd_millions']) for r in rows if 1861<=r['year']<=1869),' 1870-1900',sum(float(r['nominal_usd_millions']) for r in rows if 1870<=r['year']<=1900))
print('  Ulmer incl land 1870-1900', sum(u['c10'][str(y)]['incl_land'] for y in range(1870,1901)))

# ---------- CANALS (Cranmer 1960 Table 3 col 1, US total, millions current $; verified against page image)
can={1815:None,1816:None,1817:0.2,1818:0.7,1819:0.8,1820:1.1,1821:1.6,1822:2.7,1823:2.8,1824:2.5,1825:2.7,1826:4.0,1827:5.6,1828:7.8,1829:7.0,1830:7.5,1831:3.7,1832:4.6,1833:5.3,1834:4.4,1835:3.5,1836:4.4,1837:8.2,1838:12.3,1839:13.6,1840:14.3,1841:11.7,1842:3.1,1843:1.0,1844:1.0,1845:2.0,1846:1.8,1847:4.7,1848:4.5,1849:3.4,1850:4.9,1851:4.7,1852:3.4,1853:3.8,1854:4.7,1855:5.3,1856:4.2,1857:3.5,1858:2.8,1859:1.9,1860:1.2}
tot=sum(v for v in can.values() if v); print('canal sum',round(tot,1),'(table total 194.6)')
w('canals',[dict(year=y,nominal_usd_millions='' if v is None else fmt(v,1),note=('less than $0.05M (Cranmer "a")' if v is None else 'Cranmer (1960) Table 3 col.1, US total state+private canal investment')) for y,v in can.items()])

# ---------- ELECTRIFICATION (Ulmer 1960 Table D-1 col.3)
w('electrification',[dict(year=int(y),nominal_usd_millions=fmt(v,1),note='Ulmer (1960) Table D-1 col.3 gross capital expenditures, privately owned electric light & power utilities, excl. land') for y,v in sorted(u['d1'].items(),key=lambda kv:int(kv[0]))])

# ---------- HIGHWAYS ALL (FHWA HS 2014 DISB-C capital outlay, billions in source)
import xlrd
s=xlrd.open_workbook(R+'/fhwa_hs2014_disbc.xls').sheet_by_index(0)
rows=[]
for i in range(s.nrows):
    v=s.row_values(i)[3:]
    if isinstance(v[0],float) and 1945<=v[0]<=2014 and v[1]!='':
        rows.append(dict(year=int(v[0]),nominal_usd_millions=fmt(v[1]*1000,1),note='FHWA HS-2014 Table DISB-C capital outlay, all units of government (source sheet in $ billions, mostly 3 decimals)',
            maintenance_usd_millions=fmt(v[2]*1000,1),admin_enforcement_interest_usd_millions=fmt(v[3]*1000,1),debt_retirement_usd_millions=fmt(v[4]*1000,1)))
w('highways_all',rows,extra=('maintenance_usd_millions','admin_enforcement_interest_usd_millions','debt_retirement_usd_millions'))
